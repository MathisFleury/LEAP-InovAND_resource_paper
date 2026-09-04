#!/usr/bin/env python3.11
"""
Supplementary Excel workbook — anatomical MRI stats, group + cluster analyses.

One sheet per (comparison x measure): Autism vs NT (group) and C1/C2/C3 vs NT
(cluster, curated k-means). Long/tidy layout — one row per ROI x hemisphere,
including bilateral/midline regions. Mirrors the ROI naming used by the gt
PDF tables (16_gt_tables_anat_v2.R / 7_gt_tables_cluster.R) so region
names match across the PDF and Excel outputs.

Input:
  4_anatomical_analysis/outputs/figures/r_input_files/t_stat_anat_*.csv
  5_cluster_anatomical_analysis/outputs/tables/r_input_files/t_stat_cluster_*.csv
Output:
  _resources/Supp_Table_Anatomical_Cluster.xlsx
"""
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent.parent
GROUP_DIR = ROOT / "4_anatomical_analysis/outputs/figures/r_input_files"
CLUSTER_DIR = ROOT / "5_cluster_anatomical_analysis/outputs/tables/r_input_files"
OUT_XLSX = Path(__file__).parent / "Supp_Table_Anatomical_Cluster.xlsx"

MEASURES = [
    ("thickness", "dk", "Cortical thickness"),
    ("area", "dk", "Surface area"),
    ("grayvol", "dk", "Cortical gray matter volume"),
    ("volume", "aseg", "Subcortical volume"),
]

BASE_COLUMNS = ["ROI", "Hemisphere", "t", "p", "p_FDR", "Cohen_d",
                "d_95CI_low", "d_95CI_high", "N_group1", "N_group2"]

_PREFIX_RE = re.compile(r"^(lh_|rh_|left-|right-|Left-|Right-)")
_HEMI_L_RE = re.compile(r"^(lh_|left-|Left-)")
_HEMI_R_RE = re.compile(r"^(rh_|right-|Right-)")


def _title_hyphen_segment(s: str) -> str:
    """Mirror R's str_to_title on a hyphen-joined segment: underscore does not
    break a title-case "word" (only the very first char of the segment is
    capitalised), hyphen does."""
    return s[:1].upper() + s[1:].lower() if s else s


def roi_and_hemisphere(label: str) -> tuple[str, str]:
    if _HEMI_L_RE.match(label):
        hemi = "Left"
    elif _HEMI_R_RE.match(label):
        hemi = "Right"
    else:
        hemi = "Bilateral"
    stripped = _PREFIX_RE.sub("", label)
    roi = "-".join(_title_hyphen_segment(seg) for seg in stripped.split("-"))
    roi = roi.replace("_", " ")
    if roi.startswith("Cc "):
        roi = "CC" + roi[2:]
    return roi, hemi


def round_sig(x, sig=3):
    """Round to `sig` significant figures instead of fixed decimals, so a
    small p-value (e.g. 4e-7) survives instead of collapsing to 0.0000."""
    if pd.isna(x) or x == 0:
        return x
    from math import floor, log10
    return round(x, -int(floor(log10(abs(x)))) + (sig - 1))


def load_sheet(fp: Path) -> tuple[pd.DataFrame, list]:
    df = pd.read_csv(fp)
    has_bonf = "p_bonf" in df.columns
    columns = BASE_COLUMNS[:5] + (["p_Bonferroni"] if has_bonf else []) + BASE_COLUMNS[5:]
    df["ROI"], df["Hemisphere"] = zip(*df["label"].map(roi_and_hemisphere))
    out = df.rename(columns={
        "t_stat": "t", "p_val": "p", "p_fdr": "p_FDR", "p_bonf": "p_Bonferroni",
        "cohens_d": "Cohen_d", "cohens_d_lo": "d_95CI_low", "cohens_d_hi": "d_95CI_high",
        "n_a": "N_group1", "n_b": "N_group2",
    })[columns].copy()
    out["t"] = out["t"].round(2)
    out["p"] = out["p"].map(round_sig)
    out["p_FDR"] = out["p_FDR"].map(round_sig)
    if has_bonf:
        out["p_Bonferroni"] = out["p_Bonferroni"].map(round_sig)
    out["Cohen_d"] = out["Cohen_d"].round(3)
    out["d_95CI_low"] = out["d_95CI_low"].round(3)
    out["d_95CI_high"] = out["d_95CI_high"].round(3)
    hemi_order = {"Left": 0, "Right": 1, "Bilateral": 2}
    out = out.sort_values(by=["Hemisphere", "ROI"],
                           key=lambda c: c.map(hemi_order) if c.name == "Hemisphere" else c)
    return out.reset_index(drop=True), columns


def sheet_spec():
    for metric, atlas, label in MEASURES:
        fname = ("t_stat_anat_aseg_volume_mri_autism_vs_control.csv" if atlas == "aseg"
                  else f"t_stat_anat_dk_{metric}_mri_autism_vs_control.csv")
        yield (f"Autism_{metric}", f"{label} — Autism vs NT", GROUP_DIR / fname,
               "Group 1 = Autism, Group 2 = NT.")
    for cluster in ["C1", "C2", "C3"]:
        for metric, atlas, label in MEASURES:
            fname = f"t_stat_cluster_{cluster}_{metric}_{atlas}_mri_cluster_vs_td.csv"
            yield (f"{cluster}_{metric}", f"{label} — {cluster} vs NT", CLUSTER_DIR / fname,
                   f"Group 1 = {cluster} (Autism), Group 2 = pooled NT. Curated k-means clustering.")


CAPTION = ("Welch's two-sample t-test per region; Benjamini-Hochberg FDR correction "
           "across regions within this table. Positive Cohen's d = Group 1 > Group 2. ")

HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")


def write_sheet(wb, sheet_name, title, fp, group_note):
    if not fp.exists():
        print(f"[skip] missing {fp}")
        return
    df, columns = load_sheet(fp)
    caption = CAPTION
    if "p_Bonferroni" in columns:
        caption += "Bonferroni correction also reported. "
    ws = wb.create_sheet(sheet_name[:31])

    ws.append([title])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([caption + group_note])
    ws["A2"].font = Font(italic=True, size=9)
    ws.append([])

    header_row = 4
    ws.append(columns)
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for _, row in df.iterrows():
        ws.append(list(row))

    ws.freeze_panes = f"A{header_row + 1}"
    for i, col in enumerate(columns, start=1):
        width = max(len(col), df[col].astype(str).map(len).max() if len(df) else 0) + 2
        ws.column_dimensions[get_column_letter(i)].width = width


def main():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, title, fp, group_note in sheet_spec():
        write_sheet(wb, sheet_name, title, fp, group_note)
    wb.save(OUT_XLSX)
    print(f"Wrote {OUT_XLSX} ({len(wb.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
