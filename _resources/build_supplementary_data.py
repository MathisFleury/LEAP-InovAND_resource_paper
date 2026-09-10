#!/usr/bin/env python3.11
"""
Supplementary Data workbook (general) — gene-list reference + anatomical MRI
stats (group and cluster analyses).

Separate from build_supp_table_anatomical_cluster.py's own output (that one
stays anatomical-cluster-only); this combines the gene-list reference table
with the same anatomical/cluster sheets into one general Supplementary Data
file. Sheet 1 is the gene-level reference table (constraint metrics + curated
gene-set membership flags) shared across the paper's analyses, first since
this workbook is the general Supplementary Data file, not anatomical-only.

Input:
  /Volumes/Imaging5/EEG_MRI-MF/ALL/results/figures/supplementary_data/df_genelist.tsv
  4_anatomical_analysis/outputs/tables/r_input_files/t_stat_anat_*.csv
  5_cluster_anatomical_analysis/outputs/tables/r_input_files/t_stat_cluster_*.csv
Output:
  _resources/Supplementary_Data.xlsx
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from build_supp_table_anatomical_cluster import (
    sheet_spec, write_sheet, HEADER_FILL,
)
from openpyxl.utils import get_column_letter

GENELIST_TSV = Path("/Volumes/Imaging5/EEG_MRI-MF/ALL/results/figures/supplementary_data/df_genelist.tsv")
OUT_XLSX = Path(__file__).parent / "Supplementary_Data.xlsx"


def write_genelist_sheet(wb):
    if not GENELIST_TSV.exists():
        print(f"[skip] missing {GENELIST_TSV}")
        return
    df = pd.read_csv(GENELIST_TSV, sep="\t", low_memory=False)
    # Write booleans as literal text -- openpyxl otherwise stores a native
    # Excel boolean cell, which Excel then renders per the viewer's locale
    # (e.g. VRAI/FAUX in French) instead of the English True/False.
    bool_cols = [c for c in df.columns if df[c].dtype == bool]
    for c in bool_cols:
        df[c] = df[c].map({True: "True", False: "False"})
    ws = wb.create_sheet("Gene_List")

    ws.append(["Supplementary Data. Gene-level reference table"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["One row per gene (HGNC/Ensembl identifiers), constraint metrics "
               "(LOEUF, s_het), and curated gene-set/pathway membership flags "
               "used throughout the analyses in this resource."])
    ws["A2"].font = Font(italic=True, size=9)
    ws.append([])

    columns = list(df.columns)
    header_row = 4
    ws.append(columns)
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for _, row in df.iterrows():
        ws.append(list(row))

    ws.freeze_panes = f"A{header_row + 1}"
    for i, col in enumerate(columns, start=1):
        width = min(max(len(col), df[col].fillna("").astype(str).map(len).max() if len(df) else 0) + 2, 40)
        ws.column_dimensions[get_column_letter(i)].width = width
    print(f"  Gene_List: {len(df):,} rows x {len(columns)} cols")


def main():
    wb = Workbook()
    wb.remove(wb.active)
    write_genelist_sheet(wb)
    for sheet_name, title, fp, group_note in sheet_spec():
        write_sheet(wb, sheet_name, title, fp, group_note)
    wb.save(OUT_XLSX)
    print(f"Wrote {OUT_XLSX} ({len(wb.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
